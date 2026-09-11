from copy import deepcopy
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "PROJECT_HISTORY_AND_AI_HANDOFF.docx"


def replace_paragraph(doc, old: str, new: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text == old:
            paragraph.text = new
            return
    raise ValueError(f"Paragraph not found: {old[:80]}")


def replace_prefix(doc, prefix: str, new: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            paragraph.text = new
            return
    raise ValueError(f"Paragraph prefix not found: {prefix}")


def insert_after(paragraph, text: str) -> None:
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p.getparent().remove(new_paragraph._p)
    paragraph._p.addnext(new_paragraph._p)
    new_paragraph.style = paragraph.style
    new_paragraph.text = text


def set_cell(table, row: int, column: int, text: str) -> None:
    table.cell(row, column).text = text


def append_row_like_last(table, values: list[str]) -> None:
    if any([cell.text for cell in row.cells] == values for row in table.rows):
        return
    row_xml = deepcopy(table.rows[-1]._tr)
    table._tbl.append(row_xml)
    for cell, value in zip(table.rows[-1].cells, values, strict=True):
        cell.text = value


doc = Document(DOCX)

replace_paragraph(doc, "Asia Frontend Only Proof of Concept", "Seven Continent Frontend Only Proof of Concept")
replace_prefix(doc, "ระบบนี้คือแผนที่ความทรงจำการเดินทางส่วนบุคคล", "ระบบนี้คือแผนที่ความทรงจำการเดินทางส่วนบุคคล ไม่ใช่ trip planner และไม่ใช่แอปกรอกฟอร์มทั่วไป ผู้ใช้ควรเปิด World Atlas แล้วเลือกทวีป เห็นประวัติชีวิตการเดินทางของตนเอง จากนั้นไล่ดูตามลำดับ ทวีป ประเทศ เมือง ทริป สถานที่ รูป ค่าใช้จ่าย บันทึก และบทสรุป")
replace_paragraph(doc, "• dropdown ต้องแสดงเฉพาะประเทศในทวีปที่กำลังเปิดอยู่ โครงสร้างต้องรองรับทั้งหมด 7 ทวีป", "• dropdown ต้องแสดงเฉพาะประเทศในทวีปที่กำลังเปิดอยู่ ปัจจุบันรองรับครบทั้ง 7 ทวีป")
replace_prefix(doc, "Asia\n  Country", "World Atlas\n  Continent\n    Country\n      City\n        Trip\n          Places\n          Photos\n          Expenses\n          Notes\n          Summary")
replace_prefix(doc, "ประสบการณ์ที่ POC ต้องพิสูจน์คือ ผู้ใช้เปิดหน้า My Asia", "ประสบการณ์ที่ POC ต้องพิสูจน์คือ ผู้ใช้เปิด World Atlas สลับทวีป เห็นประเทศที่มีความทรงจำ เลือกประเทศแล้วเห็นเมืองที่เกี่ยวข้อง เปิดเมืองเพื่อดูทริปและรายละเอียด จากนั้นเพิ่มหรือแก้ไขข้อมูลได้โดยไม่สูญหายเมื่อ reload browser")
replace_paragraph(doc, "• Dashboard ระดับ Asia พร้อมสถิติและ progress", "• Dashboard ทั้ง 7 ทวีปพร้อมสถิติและ progress ที่กรองตามทวีป")
replace_paragraph(doc, "• แผนที่ 48 ประเทศจาก Static GeoJSON", "• แผนที่ 196 geographic entries จาก Static GeoJSON แยก asset ตามทวีป")
replace_prefix(doc, "แผนที่มี GeoJSON จำนวน 48 feature", "ระบบรองรับ route continent/:continentCode สำหรับ AF AN AS EU NA OC และ SA แต่ละ route โหลด catalog, GeoJSON, SVG fallback, initial camera, dropdown, trip list และสถิติของทวีปนั้น มีทั้งหมด 196 geographic entries แบ่งเป็น Asia 48, Africa 54, North America 23, South America 12, Europe 44, Australia/Oceania 14 และ Antarctica 1 ประเทศข้ามทวีปถูกจัดแบบ Asia-first: Türkiye, Cyprus, Georgia, Armenia, Azerbaijan และ Kazakhstan อยู่ Asia ส่วน Russia อยู่ Europe")
replace_prefix(doc, "Dashboard จึงแสดงจำนวนประเทศ", "Dashboard จึงแสดงจำนวนประเทศ เมือง ทริป สถานที่ วันเดินทาง ค่าใช้จ่าย และ progress โดย map เป็นองค์ประกอบหลัก ปัจจุบันตัวเลขและรายการ recent memories คำนวณจาก trip ที่อยู่ใน active continent เท่านั้น")
replace_prefix(doc, "ผู้ใช้ชี้ว่า Jump to Country", "ผู้ใช้ชี้ว่า Jump to Country ต้องแสดงเฉพาะประเทศในทวีปปัจจุบัน เพราะระบบจะมี 7 ทวีป ระบบจึงใช้ ContinentCode และฟังก์ชัน countriesByContinent Dashboard ส่งเฉพาะ catalog ของ active continent ให้ map, dropdown, form และ Smart Photo Import")
replace_prefix(doc, "โครงสร้างนี้เป็น groundwork เท่านั้น", "การขยายครบ 7 ทวีปทำด้วย dashboard เดียวและ map component เดียว ไม่ duplicate component รายทวีป แต่ใช้ metadata mapKey, center และ zoom เพื่อโหลด asset และตั้ง camera แบบ dynamic ห้ามรวมประเทศทั้งหมดไว้ใน dropdown เดียว")
replace_prefix(doc, "• เมืองใน catalog เป็นรายการ seed", "• เมืองใน catalog เป็นรายการ seed ไม่ใช่รายชื่อทุกเมืองของ 196 geographic entries คำว่า cities on map หมายถึงเมืองที่ catalog มีอยู่")
replace_prefix(doc, "• Label centroid จาก polygon หลัก", "• Label ใช้ centroid ที่ generate จาก geometry ซึ่งอาจไม่เหมาะกับประเทศที่เว้าหรือหมู่เกาะ ต้องทดสอบและ override เป็นรายกรณี")
replace_prefix(doc, "• Continent model รองรับ 7 code", "• พิกัด seed city ของ catalog นอกเอเชียใช้ country-center จากชุดข้อมูล world-countries แม้ชื่อ city จะเป็นเมืองหลวง จึงเป็นพิกัดประมาณสำหรับ POC ไม่ควรใช้เป็นข้อมูล navigation")

limitation_anchor = next(p for p in doc.paragraphs if p.text.startswith("• พิกัด seed city"))
if not any("Natural Earth 50m ไม่มี geometry ของ Tuvalu" in p.text for p in doc.paragraphs):
    insert_after(limitation_anchor, "• Antarctica ถูกแสดงผ่าน Web Mercator โดย clamp latitude ที่ -85 องศา รูปร่างจึงบิดเบือนและไม่ใช่ polar projection")
    insert_after(limitation_anchor, "• Natural Earth 50m ไม่มี geometry ของ Tuvalu ระบบจึงสร้าง diamond display geometry ขนาดเล็กเพื่อให้เลือก highlight และ focus ได้ ต้องเปลี่ยนเป็น authoritative boundary ก่อน production")

replace_paragraph(doc, "การขยายครบ 7 ทวีป", "การดูแลข้อมูลครบ 7 ทวีป")
replace_prefix(doc, "เพิ่มทีละทวีปโดยรักษา contract เดิม", "การขยายครบ 7 ทวีปเสร็จในระดับ POC แล้ว การแก้ข้อมูลต่อไปต้องรักษา contract ของ catalog, asset boundary, initial camera, country count และ test matrix ของแต่ละทวีป Dashboard ใช้ continent route และ component ชุดเดียว ห้ามสร้าง component copy แยกทวีป")
replace_prefix(doc, "เมื่อเพิ่มทวีปใหม่ให้ตรวจว่า", "เมื่อแก้ catalog ให้ตรวจว่าทุกสถิติยัง filter ตาม continent ก่อนคำนวณ ประเทศหนึ่งต้องอยู่ใน scope ที่ตกลงกันเพียงหนึ่งรายการใน catalog แม้ข้อจำกัดภูมิศาสตร์จริงจะมี transcontinental country ก็ตาม การตัดสินใจเรื่องประเทศข้ามทวีปต้องบันทึกไว้ใน catalog policy")
replace_prefix(doc, "1. เปิด My Asia Dashboard", "1. เปิด World Atlas สลับได้ทั้ง 7 ทวีป และเห็นแผนที่พร้อมชื่อประเทศตามจำนวนของแต่ละทวีป")
replace_prefix(doc, "You are continuing an Angular 22 frontend only proof of concept", "You are continuing an Angular 22 frontend only proof of concept for a one user Personal Geospatial Travel Memory Platform. The product supports all seven continents through one dynamic continent dashboard at continent/:continentCode, with 196 geographic entries and per-continent GeoJSON/SVG assets, catalog, camera, statistics and country dropdown. Visited countries and cities are always derived from Trip data. Country names remain visible for both visited and unvisited states. Visited countries use local flag SVG patterns. A selected country also shows its flag and gold outline even when unvisited, but selection must never change visited statistics. Country selection must fit the real GeoJSON geometry into the visible map area with responsive padding. SVG country paths labels and markers must stay synchronized with the MapLibre camera. Never merge all countries into one dropdown; it must remain scoped to the active continent.\n\nSmart Photo Import accepts up to 30 images, runs local Thai and English Tesseract OCR, creates a reviewable draft through PhotoAnalysisProvider, then stores structured Trip data in localStorage and image blobs in IndexedDB. No image leaves the device. The current analysis provider is heuristic, not an external LLM. Keep TravelRepository MediaStorage SummaryProvider OcrProvider and PhotoAnalysisProvider boundaries. Do not add backend database authentication cloud storage queues or client side API secrets unless the owner explicitly requests them.\n\nRead docs PROJECT_HISTORY_AND_AI_HANDOFF md and the relevant files in its Code map before making changes. Run npm run build and visually test the affected map or OCR states. Update the handoff document when a product decision or invariant changes.")
replace_paragraph(doc, "• Geographic assets ใน public/assets/maps/asia", "• Geographic assets ใน public/assets/maps/{asia,africa,north-america,south-america,europe,oceania,antarctica}")

metadata = doc.tables[0]
set_cell(metadata, 2, 1, "Personal Travel Memory และ AI Travel Atlas")
set_cell(metadata, 5, 1, "7 ทวีป รวม 196 geographic entries: Asia 48, Africa 54, North America 23, South America 12, Europe 44, Australia/Oceania 14, Antarctica 1")

code_map = doc.tables[2]
set_cell(code_map, 12, 1, "public/assets/maps/{continent}")
set_cell(code_map, 12, 2, "GeoJSON และ SVG fallback แยก 7 ทวีป")
append_row_like_last(code_map, ["13", "scripts/generate-world-maps.mjs", "สร้าง map assets และ catalog นอกเอเชียจาก world-atlas/world-countries"])
append_row_like_last(code_map, ["14", "src/app/core/data/world-countries.generated.ts", "Catalog ที่ generate สำหรับ 6 ทวีปนอกเอเชีย"])

map_tests = doc.tables[3]
set_cell(map_tests, 7, 0, "สลับทวีป")
set_cell(map_tests, 7, 1, "camera กลับภาพรวมของทวีปใหม่ map/label เปลี่ยนครบ และ dropdown มีเฉพาะประเทศของทวีปนั้น")
append_row_like_last(map_tests, ["Europe / Africa", "จำนวน feature และ dropdown ตรงกับ 44 / 54 และชื่อประเทศแสดงเสมอ"])
append_row_like_last(map_tests, ["Oceania", "ประเทศหมู่เกาะรวมถึง Tuvalu เลือก highlight และ focus ได้"])
append_row_like_last(map_tests, ["Antarctica", "geometry และ label แสดงได้ภายใต้ข้อจำกัดของ Mercator และเลือก focus ได้"])

decisions = doc.tables[5]
append_row_like_last(decisions, ["11", "Asia ใช้งานได้แล้ว ต้องการอีก 6 ทวีป", "เปลี่ยนเป็น dynamic continent route/catalog/assets/camera ชุดเดียว ครบ 7 ทวีป โดยคง continent-scoped dropdown และสถิติ"])

doc.save(DOCX)
print(DOCX)
